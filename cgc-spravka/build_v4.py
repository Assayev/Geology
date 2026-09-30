"""Справка о компании CGC (сжатая, 2 листа) на фирменном бланке из файла заказчика."""
import sys, docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, OUT = sys.argv[1], sys.argv[2]
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
d = docx.Document(SRC)
body = d.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)
FS = 11

def fmt(r, bold=False, size=FS, color=None, italic=False):
    r.bold, r.italic = bold, italic
    r.font.name = "Times New Roman"; r.font.size = Pt(size)
    r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if color: r.font.color.rgb = color
    return r

def P(text="", bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=FS, after=3, before=0, indent=None, color=None, italic=False):
    p = d.add_paragraph(); p.alignment = align
    pf = p.paragraph_format; pf.space_after = Pt(after); pf.space_before = Pt(before); pf.line_spacing = 1.0
    if indent is not None: pf.first_line_indent = Cm(indent)
    if text: fmt(p.add_run(text), bold, size, color, italic)
    return p

def H(text):
    p = P(text, True, WD_ALIGN_PARAGRAPH.LEFT, 11.5, after=3, before=7, color=NAVY)
    p.paragraph_format.keep_with_next = True

def B(lead, text):
    p = P(align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2)
    p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.first_line_indent = Cm(-0.5)
    fmt(p.add_run("• ")); fmt(p.add_run(lead), bold=True); fmt(p.add_run(text))

def shade(c):
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "E8EEF5")
    c._tc.get_or_add_tcPr().append(sh)

def T(rows, widths, head=True, bold_first_col=False, size=10):
    t = d.add_table(len(rows), len(widths)); t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 567)))
    for ri, (row, data) in enumerate(zip(t.rows, rows)):
        for ci, (c, txt, w) in enumerate(zip(row.cells, data, widths)):
            c.width = Cm(w); p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            hl = isinstance(txt, tuple)
            r = fmt(p.add_run(txt[0] if hl else txt), bold=(head and ri == 0) or (bold_first_col and ci == 0), size=size)
            if hl: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            if (head and ri == 0) or (bold_first_col and ci == 0): shade(c)
    return t

W = 16.5  # ширина набора: 21 - 3.0 - 1.5

P("Исх. № ______ от «___» ____________ 2026 г.", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, after=6)
P("СПРАВКА", True, WD_ALIGN_PARAGRAPH.CENTER, 14, after=0, color=NAVY)
P("о компании ТОО «Caspian Geology Center» (CGC)", True, WD_ALIGN_PARAGRAPH.CENTER, 12, after=4)

H("1. Общие сведения")
T([("Наименование", "ТОО «Caspian Geology Center»"),
   ("БИН / год регистрации", "230240009691 / 2023"),
   ("Юридический адрес", "010000, г. Астана, пр. Қабанбай Батыр, здание 17, нежилое помещение 15"),
   ("Производственные базы", "г. Актобе, промзона 565; Атырауская обл., Индерский р-н, п. Индербор"),
   ("Генеральный директор", "Шахтаев Г. Ж.")],
  (5.0, W - 5.0), head=False, bold_first_col=True)

H("2. Лицензии")
T([("№", "Дата", "Лицензиар", "Виды работ"),
   ("23006536", "13.03.2023", "Акимат г. Астаны (ГУ «Управление контроля и качества городской среды»)",
    "Изыскательская деятельность: инженерно-геологические и гидрогеологические работы, "
    "геофизические исследования, инженерно-геодезические работы (топосъёмка М 1:10000–1:200, "
    "съёмка подземных коммуникаций, трассирование линейных сооружений, планово-высотные сети, "
    "закладка геодезических центров)"),
   ("23007605", "31.03.2023", "Министерство энергетики РК",
    "Работы и услуги в сфере углеводородов: геофизические работы (прил. 001) и сейсморазведочные "
    "работы (прил. 002) при разведке и добыче углеводородов")],
  (1.9, 2.0, 4.1, W - 8.0))

H("3. Виды выполняемых работ")
B("Инженерно-геодезические изыскания: ", "GNSS-опорные сети, закладка реперов и уравнивание сетей; "
  "топографическая съёмка М 1:500–1:10000; съёмка подземных коммуникаций; вынос в натуру.")
B("Инженерно-геологические изыскания: ", "бурение скважин с отбором монолитов и керна, лабораторные "
  "испытания грунтов, выделение ИГЭ, графические приложения (карты, разрезы, колонки).")
B("Гидрография и батиметрия: ", "промер глубин портовых акваторий, шельфа Северного Каспия и "
  "рыбохозяйственных водоёмов; цифровые модели рельефа дна.")
B("Морская геофизика: ", "гидролокация бокового обзора, магнитометрическая съёмка, обследование "
  "дна и подводных объектов с ТНПА; геофизические и сейсморазведочные работы при разведке и добыче углеводородов.")
B("Изыскания для трубопроводов: ", "трассирование и топосъёмка трасс, инженерно-геологические "
  "изыскания по трассе; на акваториях — поиск и трассирование подводных трубопроводов "
  "(ГБО, магнитометр), батиметрия и видеообследование с ТНПА.")
B("Камеральные работы: ", "обработка данных и выпуск технических отчётов по СП РК 1.02-105-2014 "
  "и СП РК 1.02-102-2014.")

H("4. Выполненные работы (2025–2026 гг.)")
T([("№", "Заказчик", "Виды работ", "Год"),
   ("1", "ТОО «Sarzha Grain Terminal»", "Инженерно-геодезические изыскания, топосъёмка", "2026"),
   ("2", "ТОО «Sarzha Cargo Terminal»", "Инженерно-геологические и геодезические изыскания", "2026"),
   ("3", "ТОО «Sarzha Cargo Terminal»", "Инженерно-геологические изыскания, Причал №10 ММТ «Саржа» "
         "(9 скважин глубиной 14–30 м, 206 п. м)", "2026"),
   ("4", "ТОО «Sarzha Cargo Terminal»", "Закладка 11 реперов, уравнивание геодезической сети", "2026"),
   ("5", "ТОО «Semurg Invest»", "Инженерно-топографическая съёмка М 1:1000", "2026"),
   ("6", "ТОО «Semurg Invest»", "Инженерно-топографическая съёмка М 1:500", "2026"),
   ("7", "ТОО «Sarzha Cargo Terminal»", "Гидрографические работы", "2025"),
   ("8", "ТОО «Kalamkas-Khazar Operating»", "Батиметрия, Уральская седловина, Северный Каспий", "2025"),
   ("9", "ТОО «Organic Fish»", "Батиметрия акватории рыбоводных садков", "2025")],
  (0.8, 5.0, W - 7.2, 1.4))

H("5. Морское оборудование")
T([("Оборудование", "Назначение и характеристики", "Кол."),
   ("GNSS-приёмник Njord (Satlab)", "Высокоточное позиционирование судна и привязка гидрографических измерений; морское исполнение", "1"),
   ("Многолучевой эхолот NORBIT", "Батиметрическая съёмка и промер глубин, построение цифровой модели рельефа дна", "1"),
   ("Гидролокатор бокового обзора SS900", "Съёмка дна, поиск и картирование донных объектов; буксируемый, кабель-трос 50 м", "1"),
   ("Морской магнитометр SeaSPY2 (Marine Magnetics)", "Магнитная съёмка, поиск трубопроводов и ферромагнитных объектов; оверхаузеровский, "
    "чувствительность 0,01 нТл, кабель-трос 300 м", "1"),
   ("ТНПА BALTICROV BR-300", "Визуальное обследование подводных конструкций и трубопроводов; глубина до 300 м, "
    "кабель-трос 400 м, 2 видеокамеры", "1"),
   ("USBL-система Tritech Micronnav 200", "Гидроакустическое позиционирование ТНПА относительно судна", "1"),
   ("Лодка SkyBoat 460 с мотором Yamaha F60 FETL", "Работы на мелководье и в прибрежной зоне; РИБ 4,6 м, 60 л. с.; прицеп для перевозки", "1"),
   ("Генератор PATRIOT GP 3000iL, ИБП CyberPower 3000 ВА", "Автономное электропитание аппаратуры в поле и на борту", "1+1"),
   ("Спутниковый терминал Thuraya XT-LITE", "Связь вне зоны сотовых сетей", "1")],
  (5.2, W - 6.8, 1.6))

P("Специалисты Компании имеют профильное образование и опыт полевых работ на суше и акваториях "
  "Каспийского региона. Камеральная обработка ведётся на собственной вычислительной технике.", before=6, after=14, indent=1.0)

p = P(align=WD_ALIGN_PARAGRAPH.LEFT, after=0)
fmt(p.add_run("Генеральный директор\nТОО «Caspian Geology Center»"), bold=True)
p.add_run("\t" * 6); fmt(p.add_run("Шахтаев Г. Ж."), bold=True)
P("Исполнитель: Асаев Р., тел. +7 776 124 20 49, r.assayev@caspiangeo.com",
  align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5, before=14, after=0, italic=True)
d.save(OUT)
