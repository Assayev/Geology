from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin, s.right_margin = Cm(2.2), Cm(1.8)
s.top_margin, s.bottom_margin = Cm(1.2), Cm(1.8)
s.header_distance, s.footer_distance = Cm(0.8), Cm(0.7)
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.space_after = Pt(6)

def border(p, side="bottom", sz=12, color="1F3A5F"):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement(f"w:{side}")
    for k, v in (("val", "single"), ("sz", str(sz)), ("space", "4"), ("color", color)): e.set(qn(f"w:{k}"), v)
    b.append(e); pPr.append(b)

def run(p, t, bold=False, size=None, color=None, italic=False):
    r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return r

# ---- бланк: шапка ----
h = s.header; h.is_linked_to_previous = False
t = h.add_table(1, 2, Cm(17)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
c0, c1 = t.rows[0].cells; c0.width, c1.width = Cm(5.2), Cm(11.8)
c0.paragraphs[0].add_run().add_picture("cgc_logo.png", width=Cm(4.2))
p = c1.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p, "Товарищество с ограниченной ответственностью\n«Caspian Geology Center»", True, 12, NAVY)
p2 = c1.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p2, "БИН 230240009691\n010000, Республика Казахстан, г. Астана,\nпр. Қабанбай Батыр, здание 17, нежилое помещение 15", size=9.5)
hp = h.paragraphs[0]; border(hp, "bottom", 18); hp.paragraph_format.space_after = Pt(0)
# линия под таблицей: пустой абзац после таблицы
h._element.remove(hp._p); h._element.append(hp._p)

# ---- подвал ----
f = s.footer; f.is_linked_to_previous = False
fp = f.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; border(fp, "top", 8)
run(fp, "ТОО «Caspian Geology Center» · БИН 230240009691 · Производственные базы: г. Актобе, промзона 565; "
        "Атырауская обл., Индерский р-н, п. Индербор", size=8.5, color=NAVY)

# ---- тело ----
def H(text):
    p = d.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
    run(p, text, True, 12.5, NAVY)
def P(text):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25); run(p, text); return p
def B(text, lead=None):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(1.25); p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    run(p, "– "); 
    if lead: run(p, lead, True)
    run(p, text)

p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(6)
run(p, "СПРАВКА-ДОКЛАД", True, 16, NAVY)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "о компании ТОО «Caspian Geology Center» (CGC)", True, 13)
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p, "г. Астана, 2026 г.", italic=True, size=11)

H("1. Общие сведения")
P("Товарищество с ограниченной ответственностью «Caspian Geology Center» (далее – CGC, Компания) "
  "работает на рынке инженерных изысканий Республики Казахстан с 2023 года. Компания зарегистрирована "
  "в г. Астана (БИН 230240009691) и специализируется на комплексном выполнении геофизических, геодезических, "
  "геологических, гидрологических, гидрографических и иных инженерных работ для проектирования и строительства "
  "объектов промышленной, транспортной, портовой и энергетической инфраструктуры, а также для нужд "
  "углеводородной отрасли.")
tb = d.add_table(0, 2); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
rows = [("Полное наименование", "Товарищество с ограниченной ответственностью «Caspian Geology Center»"),
        ("Сокращённое наименование", "ТОО «Caspian Geology Center», CGC"),
        ("Год основания", "2023"),
        ("БИН", "230240009691"),
        ("Юридический адрес", "010000, Республика Казахстан, г. Астана, пр. Қабанбай Батыр, здание 17, нежилое помещение 15"),
        ("Производственные базы", "г. Актобе, промзона 565; Атырауская область, Индерский район, п. Индербор"),
        ("Директор", "Шахтаев Г. Ж.")]
for a, b in rows:
    r = tb.add_row().cells; r[0].width, r[1].width = Cm(5.2), Cm(11.8)
    r[0].paragraphs[0].add_run(a).bold = True; r[1].paragraphs[0].add_run(b)
    for c in r: c.paragraphs[0].paragraph_format.space_after = Pt(2)
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "E8EEF5"); r[0]._tc.get_or_add_tcPr().append(sh)

H("2. Направления деятельности")
P("Компания выполняет полный цикл инженерных изысканий – от полевых работ до камеральной обработки и выпуска "
  "отчётной документации:")
B(" полевые исследования грунтов, бурение и опробование, лабораторные испытания, "
  "литологическое расчленение разреза, оценка физико-механических свойств грунтов;", "инженерно-геологические работы –")
B(" изучение подземных вод, определение уровней и режима, гидрогеологические исследования "
  "площадок и трасс;", "гидрогеологические и гидрологические работы –")
B(" рекогносцировка и съёмка, сейсморазведочные работы, исследования для инженерных изысканий "
  "и для разведки и добычи углеводородов;", "геофизические работы –")
B(" топографические съёмки масштабов от 1:10000 до 1:200, съёмка подземных коммуникаций и "
  "сооружений, трассирование и съёмка линейных объектов, создание планово-высотных съёмочных сетей, "
  "построение и закладка геодезических центров, вынос в натуру и привязка выработок и точек изысканий;", "инженерно-геодезические работы –")
B(" промеры глубин, батиметрическая съёмка акваторий, построение цифровых моделей дна "
  "для проектирования морских, речных и портовых сооружений;", "гидрографические работы –")
B(" подготовка технических отчётов, ведомостей, картографических и графических материалов "
  "(карты фактического материала, геолого-литологические разрезы и колонки).", "камеральная обработка и отчётность –")

H("3. Разрешительная документация")
P("Деятельность Компании осуществляется на основании государственных лицензий (неотчуждаемые, класс 1):")
lt = d.add_table(1, 4); lt.style = "Table Grid"; lt.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, txt in zip(lt.rows[0].cells, ("№ лицензии", "Дата выдачи", "Вид деятельности", "Лицензиар")):
    c.paragraphs[0].add_run(txt).bold = True
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), "E8EEF5"); c._tc.get_or_add_tcPr().append(sh)
lic = [("23006536", "13.03.2023\n(первичная – 02.03.2023)",
        "Изыскательская деятельность: инженерно-геологические и инженерно-гидрогеологические работы; "
        "геофизические исследования; инженерно-геодезические работы",
        "ГУ «Управление контроля и качества городской среды г. Астаны», Акимат г. Астаны"),
       ("23007605", "31.03.2023",
        "Работы и услуги в сфере углеводородов: геофизические работы (прил. 001 от 31.03.2023) и "
        "сейсморазведочные работы (прил. 002 от 31.05.2023) при разведке и добыче углеводородов",
        "Министерство энергетики Республики Казахстан")]
for row in lic:
    cs = lt.add_row().cells
    for c, txt in zip(cs, row): c.paragraphs[0].add_run(txt).font.size = Pt(10.5)
for r in lt.rows:
    for c, w in zip(r.cells, (2.2, 3.0, 7.3, 4.5)): c.width = Cm(w)
    for c in r.cells: c.paragraphs[0].paragraph_format.space_after = Pt(2)

H("4. Реализованные проекты")
P("Среди выполненных Компанией работ – инженерные изыскания по объекту «Многофункциональный морской терминал "
  "«Саржа». Универсальный терминал. Причал №10 с прилегающей площадкой» (заказчик – ТОО «Sarzha Cargo Terminal», "
  "г. Актау), включающие инженерно-геологические исследования, составление карты фактического материала, "
  "геолого-литологических разрезов и колонок и выпуск технического отчёта в 2026 году.")

H("5. Преимущества Компании")
B(" все виды изысканий – геологические, геодезические, геофизические, гидрологические, гидрографические – "
  "выполняются одной организацией, что сокращает сроки и исключает разрыв ответственности;", "комплексность:")
B(" действующие лицензии на изыскательскую деятельность и на работы в сфере углеводородов;", "правовая обеспеченность:")
B(" производственные базы в Актюбинской и Атырауской областях, работа в Каспийском регионе;", "региональное присутствие:")
B(" отчётные материалы оформляются в соответствии с требованиями заказчиков и действующих нормативных документов РК.", "качество документации:")

H("6. Заключение")
P("ТОО «Caspian Geology Center» обладает необходимыми разрешительными документами, производственной базой и "
  "профильной специализацией для выполнения инженерных изысканий любой сложности и готово к сотрудничеству "
  "с проектными, строительными и недропользующими организациями.")

p = d.add_paragraph(); p.paragraph_format.space_before = Pt(22)
run(p, "Директор\t\t\t\t\t\t\tШахтаев Г. Ж.", True)
d.save("Справка_о_компании_CGC.docx")
