"""Build documents on the CGC (Caspian Geology Center) letterhead.

The template docx is the client's original file, untouched: the letterhead
(logo, 3-language address table, dividers, 5% watermark) lives in its header,
page size / margins live in its sectPr. We only replace the body.

    from cgc_doc import new_doc, para, table, save
    doc = new_doc()
    para(doc, "Текст", bold=True, align="center")
    save(doc, "out.docx")          # + out.pdf via soffice if pdf=True
"""
import copy, os, subprocess, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "assets", "CGC_letterhead_template.docx")
BODY_FONT, BODY_SIZE = "Times New Roman", 12
ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
         "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def new_doc(font=BODY_FONT, size=BODY_SIZE):
    doc = Document(TEMPLATE)
    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):      # drop the template's empty paragraph
        body.remove(p)                            # (sectPr with header ref/margins stays)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = font, Pt(size)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), font)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rf.attrib.pop(qn(a), None)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0
    return doc


def para(doc, text="", bold=False, italic=False, align="left", size=None,
         after=6, before=0, first_indent_cm=None, left_indent_cm=None):
    p = doc.add_paragraph()
    f = p.paragraph_format
    p.alignment = ALIGN[align]
    f.space_after, f.space_before = Pt(after), Pt(before)
    if first_indent_cm is not None: f.first_line_indent = Cm(first_indent_cm)
    if left_indent_cm is not None: f.left_indent = Cm(left_indent_cm)
    if text:
        r = p.add_run(text); r.bold, r.italic = bold, italic
        if size: r.font.size = Pt(size)
    return p


def table(doc, rows, widths_cm=None, header=True, size=None):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            r.bold = header and i == 0
            if size: r.font.size = Pt(size)
            if widths_cm: c.width = Cm(widths_cm[j])
    return t


def _restore_header(path):
    """python-docx re-serialises header parts; put back the client's original bytes."""
    import zipfile, shutil
    keep = ("word/header1.xml", "word/_rels/header1.xml.rels")
    src = zipfile.ZipFile(TEMPLATE)
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            zout.writestr(it, src.read(it.filename) if it.filename in keep else zin.read(it.filename))
    shutil.move(tmp, path)


def _render_copy(path, tmp):
    """LibreOffice ignores <a:alphaModFix> (5% watermark) and draws 'outset' borders
    too thick. Word does both right, so the .docx stays untouched; only the copy that
    is fed to LibreOffice gets the watermark pre-faded and the dividers set to 0.75pt."""
    import zipfile, io
    from PIL import Image
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "word/media/image1.png":
                im = Image.open(io.BytesIO(data)).convert("RGBA")
                im.putalpha(im.getchannel("A").point(lambda a: round(a * 0.05)))
                b = io.BytesIO(); im.save(b, "PNG"); data = b.getvalue()
            elif it.filename == "word/header1.xml":
                data = data.replace(b'<w:insideV w:val="outset" w:sz="6"', b'<w:insideV w:val="single" w:sz="6"')
                data = data.replace(b'<a:alphaModFix amt="5000"/>', b"")
            zout.writestr(it, data)


def save(doc, path, pdf=False):
    """Save docx (letterhead byte-identical to the template). pdf=True also renders
    a PDF next to it via LibreOffice (needs libreoffice-writer + fonts-liberation)."""
    doc.save(path)
    _restore_header(path)
    if pdf:
        import tempfile
        outdir = os.path.dirname(os.path.abspath(path))
        with tempfile.TemporaryDirectory() as td:
            tmp = os.path.join(td, os.path.basename(path))
            _render_copy(path, tmp)
            subprocess.run(["soffice", "--headless", "-env:UserInstallation=file:///tmp/lo_prof",
                            "--convert-to", "pdf", "--outdir", outdir, tmp],
                           check=True, capture_output=True)
    return path
