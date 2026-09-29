import docx, copy, sys
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
W='/tmp/claude-0/-home-user-Geology/af002989-109f-5670-bf4b-111a3c9940db/scratchpad/w/'
ETH=W+'38c25f9e-Business_Ethics_Compliance_Certificate_Contractor_or_Suppliers_EN-RU.docx'
VQD=W+'9eebeb7a-VQD_-_Vendor_Qualification_Undertaking_EN-RU_combined1.docx'

def set_text(p, text):
    rs=p.runs
    rs[0].text=text
    for r in rs[1:]: r._r.getparent().remove(r._r)

def replace_in_par(p, old, new):
    full=''.join(r.text for r in p.runs)
    if old not in full: return False
    # keep first-run formatting
    set_text(p, full.replace(old,new)); return True

def letterhead(doc, logo_w=Cm(3.6)):
    for s in doc.sections:
        h=s.header; h.is_linked_to_previous=False
        ps=h.paragraphs
        for p in ps[1:]: p._p.getparent().remove(p._p)
        p=ps[0]
        for r in list(p.runs): r._r.getparent().remove(r._r)
        p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run().add_picture('cgc_logo.png', width=logo_w)
        s.top_margin=Cm(2.2); s.header_distance=Cm(1.0)

def add_par_after(p, text='', bold=False):
    new=copy.deepcopy(p._p); p._p.addnext(new)
    np_=docx.text.paragraph.Paragraph(new,p._parent)
    for r in list(np_.runs): r._r.getparent().remove(r._r)
    if text: 
        r=np_.add_run(text); r.bold=bold
    return np_

# ---------- Ethics certificate
d=docx.Document(ETH)
letterhead(d)
for p in d.paragraphs:
    replace_in_par(p,'[name], [office or title] of [Contractor]','G. Zh. Shakhtayev, Director of Caspian Geology Center LLP')
    replace_in_par(p,'[Contractor] (the “Contractor”)','Caspian Geology Center LLP (the “Contractor”)')
    replace_in_par(p,'[название], [офис или должность][Подрядчик]','Шахтаевым Г. Ж., Директором ТОО «Caspian Geology Center»')
    if p.text.startswith('Children, spouses'):
        replace_in_par(p,' On behalf of [Contractor]','')
        sig=p
# signature block after "Date: ..." paragraph -> build before it
datep=[p for p in d.paragraphs if p.text.startswith('Date:')][0]
tbl=d.add_table(rows=4,cols=2)
rows=[('On behalf of Caspian Geology Center LLP',''),
      ('Signature:','…………………………………………'),
      ('Name / Title:','G. Zh. Shakhtayev, Director'),
      ('Stamp:','')]
for r,(a,b) in zip(tbl.rows,rows):
    r.cells[0].text=a; r.cells[1].text=b
    for c in r.cells:
        for pp in c.paragraphs:
            pp.paragraph_format.space_after=Pt(6)
            for rr in pp.runs: rr.font.size=Pt(11)
    r.cells[0].width=Cm(6); r.cells[1].width=Cm(10)
for rr in tbl.rows[0].cells[0].paragraphs[0].runs: rr.bold=True
datep._p.addprevious(tbl._tbl)
for r in tbl.rows:
    for c in r.cells:
        for pp in c.paragraphs: pp.paragraph_format.keep_with_next=True
datep.paragraph_format.keep_with_next=False
# убрать пустые абзацы между текстом и блоком подписи
prev=tbl._tbl.getprevious()
while prev is not None and prev.tag.endswith('}p') and not ''.join(prev.itertext()).strip():
    x=prev.getprevious(); prev.getparent().remove(prev); prev=x

ru=[p for p in d.paragraphs if p.text.startswith('Русская версия')][0]
nx=datep._p.getnext()
while nx is not None and nx is not ru._p:
    n2=nx.getnext()
    if nx.tag.endswith('}p') and not ''.join(nx.itertext()).strip(): nx.getparent().remove(nx)
    nx=n2
ru.paragraph_format.page_break_before=True
d.save('Business_Ethics_Compliance_Certificate_CGC.docx')

# ---------- VQD
d=docx.Document(VQD)
letterhead(d)
t=d.tables[0]
def cell(r,c,text):
    ce=t.rows[r].cells[c]; ps=ce.paragraphs
    if ps[0].runs: set_text(ps[0],text)
    else: ps[0].add_run(text)
cell(3,1,'G. Zh. Shakhtayev')
t.rows[4].cells[0].paragraphs[0].add_run('Title:'); t.rows[4].cells[1].paragraphs[0].add_run('Director, Caspian Geology Center LLP')
cell(6,1,'…………………………, Astana, Kazakhstan')
cell(9,1,'[stamp]') if False else None
d.save('Vendor_Qualification_Undertaking_CGC.docx')
